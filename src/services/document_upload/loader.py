import logging
from pathlib import Path
from typing import Any, Dict, List, Tuple

from src.config import UploadSettings
from src.exceptions import (
    UnsupportedUploadFormatError,
    UploadFileTooLargeError,
    UploadParsingException,
    UploadValidationError,
)
from src.schemas.pdf_parser.models import PdfContent

from .models import UploadParserType

logger = logging.getLogger(__name__)


def validate_upload_file(file_path: Path, settings: UploadSettings) -> None:
    """校验文件是否存在、扩展名白名单及大小限制。"""
    if not file_path.exists():
        raise UploadValidationError(f"File not found: {file_path}")

    if not file_path.is_file():
        raise UploadValidationError(f"Path is not a file: {file_path}")

    extension = file_path.suffix.lower()
    if extension not in settings.allowed_extension_set():
        allowed = ", ".join(sorted(settings.allowed_extension_set()))
        raise UnsupportedUploadFormatError(
            f"Unsupported file type '{extension}'. Allowed extensions: {allowed}"
        )

    size_mb = file_path.stat().st_size / (1024 * 1024)
    if size_mb > settings.max_file_size_mb:
        raise UploadFileTooLargeError(
            f"File size {size_mb:.2f}MB exceeds limit of {settings.max_file_size_mb}MB"
        )


def derive_title_hint(file_path: Path, override: str | None = None) -> str:
    if override and override.strip():
        return override.strip()
    return file_path.stem.replace("_", " ").replace("-", " ").strip() or file_path.name


def _ensure_non_empty_text(text: str, file_path: Path, settings: UploadSettings) -> str:
    cleaned = text.strip()
    if len(cleaned) < settings.min_extracted_text_chars:
        raise UploadParsingException(
            f"Extracted text from '{file_path.name}' is too short "
            f"({len(cleaned)} chars, minimum {settings.min_extracted_text_chars})"
        )
    return cleaned


def _pdf_sections_to_dicts(pdf_content: PdfContent) -> List[Dict[str, str]]:
    return [{"title": section.title, "content": section.content} for section in pdf_content.sections]


def parse_pdf_content(
    pdf_content: PdfContent,
    file_path: Path,
    settings: UploadSettings,
    title_hint: str | None = None,
) -> Tuple[str, List[Dict[str, str]] | None, Dict[str, Any]]:
    raw_text = _ensure_non_empty_text(pdf_content.raw_text, file_path, settings)
    sections = _pdf_sections_to_dicts(pdf_content) or None
    metadata = {
        "parser_metadata": pdf_content.metadata or {},
        "figures_count": len(pdf_content.figures),
        "tables_count": len(pdf_content.tables),
        "references_count": len(pdf_content.references),
        "title_hint": derive_title_hint(file_path, title_hint),
    }
    return raw_text, sections, metadata


def extract_text_file(file_path: Path, settings: UploadSettings, parser_type: UploadParserType) -> str:
    try:
        text = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = file_path.read_text(encoding="utf-8", errors="replace")
    return _ensure_non_empty_text(text, file_path, settings)


def extract_docx_text(file_path: Path, settings: UploadSettings) -> Tuple[str, Dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise UploadParsingException("python-docx is required to parse .docx files") from exc

    document = Document(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n".join(paragraphs)
    text = _ensure_non_empty_text(text, file_path, settings)

    title_hint = derive_title_hint(file_path)
    try:
        core_title = document.core_properties.title
        if core_title and core_title.strip():
            title_hint = core_title.strip()
    except Exception:
        logger.debug("Could not read DOCX core properties title for %s", file_path.name)

    metadata = {"paragraph_count": len(paragraphs), "title_hint": title_hint}
    return text, metadata


def extract_excel_text(file_path: Path, settings: UploadSettings) -> Tuple[str, Dict[str, Any]]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise UploadParsingException("pandas is required to parse Excel files") from exc

    try:
        excel_file = pd.ExcelFile(file_path)
    except Exception as exc:
        raise UploadParsingException(f"Failed to open Excel file '{file_path.name}': {exc}") from exc

    parts: List[str] = []
    sheet_stats: Dict[str, int] = {}

    for sheet_name in excel_file.sheet_names:
        dataframe = excel_file.parse(sheet_name)
        if dataframe.empty:
            continue

        truncated = dataframe.head(settings.excel_max_rows)
        parts.append(f"工作表: {sheet_name}\n{truncated.to_string(index=False)}")
        sheet_stats[sheet_name] = len(truncated)

    if not parts:
        raise UploadParsingException(f"No readable sheet content found in '{file_path.name}'")

    text = _ensure_non_empty_text("\n\n".join(parts), file_path, settings)
    metadata = {
        "sheet_count": len(excel_file.sheet_names),
        "sheet_row_counts": sheet_stats,
        "excel_max_rows": settings.excel_max_rows,
        "title_hint": derive_title_hint(file_path),
    }
    return text, metadata


def extract_by_extension(
    file_path: Path,
    settings: UploadSettings,
    title_hint: str | None = None,
) -> Tuple[str, UploadParserType, List[Dict[str, str]] | None, Dict[str, Any]]:
    """同步提取非 PDF 格式文本；PDF 由 service 层通过 Docling 处理。"""
    extension = file_path.suffix.lower()
    resolved_title = derive_title_hint(file_path, title_hint)

    if extension in {".txt"}:
        text = extract_text_file(file_path, settings, UploadParserType.TXT)
        return text, UploadParserType.TXT, None, {"title_hint": resolved_title}

    if extension in {".md"}:
        text = extract_text_file(file_path, settings, UploadParserType.MARKDOWN)
        return text, UploadParserType.MARKDOWN, None, {"title_hint": resolved_title}

    if extension == ".docx":
        text, metadata = extract_docx_text(file_path, settings)
        return text, UploadParserType.DOCX, None, metadata

    if extension in {".xlsx", ".xls"}:
        text, metadata = extract_excel_text(file_path, settings)
        return text, UploadParserType.EXCEL, None, metadata

    allowed = ", ".join(sorted(settings.allowed_extension_set()))
    raise UnsupportedUploadFormatError(
        f"Unsupported file type '{extension}'. Allowed extensions: {allowed}"
    )
