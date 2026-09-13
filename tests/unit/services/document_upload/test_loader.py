import pytest
from docx import Document

from src.config import UploadSettings
from src.exceptions import (
    UnsupportedUploadFormatError,
    UploadFileTooLargeError,
    UploadParsingException,
)
from src.services.document_upload.loader import (
    derive_title_hint,
    extract_by_extension,
    extract_docx_text,
    extract_excel_text,
    extract_text_file,
    validate_upload_file,
)
from src.services.document_upload.models import UploadParserType


@pytest.fixture
def upload_settings(tmp_path):
    return UploadSettings(
        upload_dir=str(tmp_path / "uploads"),
        max_file_size_mb=1,
        allowed_extensions=[".pdf", ".txt", ".md", ".docx", ".xlsx", ".xls"],
        excel_max_rows=10,
        min_extracted_text_chars=10,
    )


def test_upload_settings_resolved_dir(upload_settings, tmp_path):
    resolved = upload_settings.resolved_upload_dir()
    assert resolved.exists()
    assert resolved.name == "uploads"


def test_validate_upload_file_rejects_unknown_extension(upload_settings, tmp_path):
    bad_file = tmp_path / "sample.exe"
    bad_file.write_bytes(b"binary")

    with pytest.raises(UnsupportedUploadFormatError):
        validate_upload_file(bad_file, upload_settings)


def test_validate_upload_file_rejects_large_file(upload_settings, tmp_path):
    large_file = tmp_path / "large.txt"
    large_file.write_text("x" * (2 * 1024 * 1024))

    with pytest.raises(UploadFileTooLargeError):
        validate_upload_file(large_file, upload_settings)


def test_extract_text_file(upload_settings, tmp_path):
    text_file = tmp_path / "notes.txt"
    text_file.write_text("这是一段用于测试上传解析的中文文本内容。", encoding="utf-8")

    text = extract_text_file(text_file, upload_settings, UploadParserType.TXT)
    assert "测试上传解析" in text


def test_extract_markdown_file(upload_settings, tmp_path):
    md_file = tmp_path / "notes.md"
    md_file.write_text("# Title\n\nMarkdown body for upload parsing test.", encoding="utf-8")

    text, parser_used, sections, metadata = extract_by_extension(md_file, upload_settings)
    assert parser_used == UploadParserType.MARKDOWN
    assert "Markdown body" in text
    assert sections is None
    assert metadata["title_hint"] == "notes"


def test_extract_docx_file(upload_settings, tmp_path):
    docx_file = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("第一段上传文档测试内容。")
    document.add_paragraph("第二段上传文档测试内容。")
    document.core_properties.title = "测试报告"
    document.save(docx_file)

    text, metadata = extract_docx_text(docx_file, upload_settings)
    assert "第一段上传文档测试内容" in text
    assert metadata["title_hint"] == "测试报告"


def test_extract_excel_file(upload_settings, tmp_path):
    pytest.importorskip("openpyxl")
    import pandas as pd

    excel_file = tmp_path / "table.xlsx"
    dataframe = pd.DataFrame({"name": ["Alice", "Bob"], "score": [90, 88]})
    with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="Sheet1", index=False)

    text, metadata = extract_excel_text(excel_file, upload_settings)
    assert "Alice" in text
    assert metadata["sheet_count"] == 1


def test_extract_by_extension_rejects_short_text(upload_settings, tmp_path):
    text_file = tmp_path / "short.txt"
    text_file.write_text("too short", encoding="utf-8")

    with pytest.raises(UploadParsingException):
        extract_by_extension(text_file, upload_settings)


def test_derive_title_hint_prefers_override(tmp_path):
    file_path = tmp_path / "raw-name.pdf"
    file_path.write_bytes(b"%PDF-1.4")

    assert derive_title_hint(file_path, "自定义标题") == "自定义标题"
